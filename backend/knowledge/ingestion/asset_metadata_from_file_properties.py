"""Phase 6A.3 corrective addendum: an OPTIONAL, read-only extractor that
populates a `KnowledgeAssetMetadata` instance from a real DOCX/XLSX
file's own OOXML core/document properties -- one of the legitimate
metadata sources §6 names ("file properties", "filename").

DELIBERATELY NOT WIRED INTO `local_file_adapter.py`'S DEFAULT INGESTION
PATH: this addendum's own instruction is explicit -- "Do NOT reopen or
redesign the 6A.3 multimodal ingestion architecture." `ingest_local_file`/
`ingest_and_structure_local_file(s)` (6A.3) are completely unmodified by
this module; a caller who wants asset metadata populated from file
properties must explicitly call the functions below and merge the result
into an `IngestedKnowledgeDocument.metadata.asset_metadata` themselves.
This keeps 6A.3's own frozen output shape byte-for-byte unchanged for
every existing caller.

READ-ONLY, NO SOURCE MODIFICATION: both functions only ever call
`docx.Document(path)`/`openpyxl.load_workbook(path)` in read mode and
read attributes off the resulting `core_properties`/`properties` object
-- neither function writes to `path`, and neither ever opens the file in
write mode. Never modifies, moves, or deletes the original file (§27).

DELIBERATE NON-MAPPINGS (documented, not omissions):
- OOXML's own `core_properties.revision`/`properties.revision` is a
  bare integer SAVE COUNT (how many times the file has been saved in
  the authoring application) -- NOT a document revision LABEL like
  "PA1". Mapping it to `KnowledgeAssetIdentityMetadata`'s conceptual
  "Revision" (itself already represented by `KnowledgeVersion.revision`,
  not duplicated here -- see asset_metadata.py's own docstring) would
  misrepresent a save-count as a formal revision identifier. This
  function does NOT read or expose that property at all.
- `core_properties.title`/`properties.title` is NOT mapped anywhere --
  `KnowledgeObject.title` already has an authoritative source elsewhere
  in the ingestion pipeline (the document's own extracted heading/
  filename); duplicating it here from a possibly-stale Office metadata
  field would risk two disagreeing titles for the same object.
- Confidentiality classification is NEVER derived from file properties
  here -- no OOXML core property reliably or deterministically encodes
  it; attempting to infer it from `category`/`keywords`/`comments` text
  would require heuristic pattern-matching against banner text, which
  this addendum's own "no metadata guessing" discipline forbids.
- `identity.date` ("date of last document update" as the document ITSELF
  claims, typically a header/body fact) is NOT populated from file
  properties here -- `core_properties.modified`/`properties.modified` is
  a FILE-SYSTEM-level fact, mapped instead to `audit.last_modified_date`
  (see `KnowledgeAssetAuditMetadata`'s own docstring for why the two are
  kept distinct).

Every populated field's `source` is `MetadataSource.FILE_PROPERTIES`
(for a true OOXML core-property value) or `MetadataSource.FILENAME` (for
`file_format`, derived from the path's own suffix, never from file
content) -- never guessed, never blank-becomes-a-default.
"""
from __future__ import annotations

from pathlib import Path

from backend.knowledge.domain.asset_metadata import (
    DateMetadataField,
    KnowledgeAssetMetadata,
    MetadataSource,
    TextMetadataField,
    normalize_date_value,
)

__all__ = ["extract_asset_metadata_from_docx_properties", "extract_asset_metadata_from_xlsx_properties"]


def _file_format_field(path: Path) -> TextMetadataField:
    suffix = path.suffix.lstrip(".").lower()
    if not suffix:
        return TextMetadataField()
    return TextMetadataField(raw_value=suffix, normalized_value=suffix, source=MetadataSource.FILENAME)


def _text_field(value: object) -> TextMetadataField:
    if not value:
        return TextMetadataField()
    text = str(value).strip()
    if not text:
        return TextMetadataField()
    return TextMetadataField(raw_value=text, source=MetadataSource.FILE_PROPERTIES)


def _date_field(value: object) -> DateMetadataField:
    if value is None:
        return DateMetadataField()
    # python-docx/openpyxl already return a real `datetime`/`date` object
    # for `created`/`modified` when the property is present -- format it
    # as ISO-8601 so `normalize_date_value` (which expects a string, the
    # same contract every other caller of this function uses) resolves
    # it deterministically rather than adding a second, object-typed
    # code path.
    raw = value.isoformat() if hasattr(value, "isoformat") else str(value)
    return DateMetadataField(
        raw_value=raw,
        normalized_value=normalize_date_value(raw),
        source=MetadataSource.FILE_PROPERTIES,
    )


def extract_asset_metadata_from_docx_properties(path: Path) -> KnowledgeAssetMetadata:
    """Read-only extraction from a real `.docx` file's own OOXML core
    properties. Raises whatever `python-docx` itself raises for an
    unreadable/corrupt/non-DOCX file -- this function performs no
    fallback/guessing of its own for a failed read.
    """
    import docx  # noqa: PLC0415 -- imported lazily, matching this codebase's "only import a heavy dependency where used" convention.

    document = docx.Document(str(path))
    core = document.core_properties

    metadata = KnowledgeAssetMetadata()
    metadata.identity.file_format = _file_format_field(path)
    metadata.identity.language_code = _text_field(core.language)
    metadata.roles.prepared_by = _text_field(core.author)
    metadata.audit.last_modified_by = _text_field(core.last_modified_by)
    metadata.audit.last_modified_date = _date_field(core.modified)
    return metadata


def extract_asset_metadata_from_xlsx_properties(path: Path) -> KnowledgeAssetMetadata:
    """Read-only extraction from a real `.xlsx` file's own OOXML
    workbook properties. Raises whatever `openpyxl` itself raises for an
    unreadable/corrupt/non-XLSX file. Uses `read_only=True` deliberately
    -- unlike `backend/knowledge/ingestion/extractors/xlsx.py`'s own
    6A.3 extractor (which needs `read_only=False` for `.tables`/
    `.dimensions`), this function only ever reads `workbook.properties`,
    which IS available in read-only mode -- confirmed by direct
    interactive check before writing this function, mirroring the same
    empirical-first discipline the 6A.3 XLSX extractor change used.
    """
    import openpyxl  # noqa: PLC0415 -- imported lazily, matching this codebase's "only import a heavy dependency where used" convention.

    workbook = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    try:
        props = workbook.properties
        metadata = KnowledgeAssetMetadata()
        metadata.identity.file_format = _file_format_field(path)
        metadata.identity.language_code = _text_field(props.language)
        metadata.roles.prepared_by = _text_field(props.creator)
        metadata.audit.last_modified_by = _text_field(props.lastModifiedBy)
        metadata.audit.last_modified_date = _date_field(props.modified)
        return metadata
    finally:
        workbook.close()
