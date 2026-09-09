"""XLSX workbook extraction (A5 Layer D).

Preserves structure as Workbook -> Sheet -> header/row schema -- never
flattened into one arbitrary text blob (A5 instruction section 23).
Formulas are read as their literal string form (`data_only=False`,
openpyxl's default) and NEVER evaluated; no macro (`xlsm`) content is
ever executed (openpyxl has no VBA execution capability at all -- macro
code, if present, is opaque binary this module never touches). Row
sampling is bounded (`_MAX_SAMPLE_ROWS`) so one enormous sheet cannot
blow up a single artifact's `extracted_text` size -- the real row count
is still recorded so a truncated sample is never presented as complete.

An empty template (headers only, zero data rows) still produces a real,
non-degenerate sheet artifact -- the header/schema row alone is useful
knowledge (A5 instruction section 23), never treated as "nothing to
extract."
"""
from __future__ import annotations

import io
from typing import Optional

from backend.knowledge.domain.artifacts import ArtifactExtractionStatus, KnowledgeArtifact
from backend.knowledge.ingestion.extraction import ExtractionBudget, deterministic_artifact_id, hash_bytes

_MAX_SAMPLE_ROWS = 25


def _render_row(row: tuple) -> str:
    return " | ".join("" if cell is None else str(cell) for cell in row)


def _render_sheet(sheet_name: str, rows: list[tuple], total_row_count: int) -> str:
    lines = [f"Sheet: {sheet_name}"]
    if not rows:
        lines.append("(no rows)")
        return "\n".join(lines)

    header, *data_rows = rows
    lines.append("Headers: " + _render_row(header))
    for row in data_rows:
        lines.append(_render_row(row))
    if total_row_count > len(rows):
        lines.append(f"... ({total_row_count - len(rows)} more row(s) not shown)")
    return "\n".join(lines)


def extract_xlsx(data: bytes, *, container_artifact_id: str, depth: int, budget: ExtractionBudget) -> tuple[str, list[KnowledgeArtifact]]:
    """Returns (workbook summary text, one `KnowledgeArtifact` per sheet,
    each with `parent_artifact_id=container_artifact_id`,
    `kind="xlsx_sheet"`, `depth=depth`, and `locator_detail="sheet=<name>"`).
    """
    import openpyxl  # noqa: PLC0415 -- imported lazily, matching this codebase's "only import a heavy dependency where used" convention.

    workbook = openpyxl.load_workbook(io.BytesIO(data), data_only=False, read_only=True)

    sheet_artifacts: list[KnowledgeArtifact] = []
    sheet_names: list[str] = list(workbook.sheetnames)

    for position, sheet_name in enumerate(sheet_names):
        worksheet = workbook[sheet_name]
        rows: list[tuple] = []
        total_row_count = 0
        for row in worksheet.iter_rows(values_only=True):
            total_row_count += 1
            if len(rows) < _MAX_SAMPLE_ROWS:
                rows.append(row)

        sheet_text = _render_sheet(sheet_name, rows, total_row_count)
        sheet_hash = hash_bytes(sheet_text.encode("utf-8"))

        try:
            budget.reserve_bytes(len(sheet_text.encode("utf-8")))
            budget.reserve_artifact_slot()
        except Exception:
            budget.record_skipped(f"{sheet_name} (sheet)", "extraction limit exceeded")
            continue

        artifact_id = deterministic_artifact_id(
            parent_artifact_id=container_artifact_id, kind="xlsx_sheet", position=position, content_hash=sheet_hash
        )
        sheet_artifacts.append(
            KnowledgeArtifact(
                artifact_id=artifact_id,
                parent_artifact_id=container_artifact_id,
                kind="xlsx_sheet",
                media_type="text/plain",
                display_name=sheet_name,
                depth=depth,
                content_hash=sheet_hash,
                extracted_text=sheet_text,
                derived=False,
                locator_detail=f"sheet={sheet_name}",
                extraction_status=ArtifactExtractionStatus.COMPLETE,
            )
        )
        budget.record_artifact(len(sheet_text.encode("utf-8")))

    workbook.close()
    summary = f"Workbook with {len(sheet_names)} sheet(s): " + ", ".join(sheet_names)
    return summary, sheet_artifacts
