"""Phase 6A.3 corrective addendum: the optional, read-only file-properties
extractor (`backend/knowledge/ingestion/asset_metadata_from_file_properties.py`).

Covers: deterministic extraction from real OOXML core properties, no
source-file modification (hash-before/after proof), the deliberate
non-mapping of OOXML's numeric `revision` (save count, not a document
revision label), and safe behavior when a property is entirely absent.
"""
from __future__ import annotations

import ast
import hashlib
import inspect
from datetime import datetime, timezone
from pathlib import Path

import pytest

from backend.knowledge.domain.asset_metadata import MetadataSource
from backend.knowledge.ingestion import asset_metadata_from_file_properties as _module_under_test
from backend.knowledge.ingestion.asset_metadata_from_file_properties import (
    extract_asset_metadata_from_docx_properties,
    extract_asset_metadata_from_xlsx_properties,
)


def _sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _revision_attribute_accesses_in_module() -> list[str]:
    """AST-based (not docstring-substring-based) proof that no executable
    code in this module ever reads a `.revision` attribute off anything
    -- deliberately excludes docstrings/comments, which legitimately
    discuss the non-mapping in prose.
    """
    tree = ast.parse(inspect.getsource(_module_under_test))
    return [
        ast.dump(node)
        for node in ast.walk(tree)
        if isinstance(node, ast.Attribute) and node.attr == "revision"
    ]


@pytest.fixture()
def synthetic_docx(tmp_path: Path) -> Path:
    from docx import Document

    document = Document()
    document.add_heading("Synthetic Procedure", level=1)
    document.add_paragraph("Invented, non-sensitive body text.")
    document.core_properties.author = "Synthetic Author"
    document.core_properties.language = "en-US"
    document.core_properties.last_modified_by = "Synthetic Modifier"
    document.core_properties.modified = datetime(2026, 8, 14, 10, 0, 0, tzinfo=timezone.utc)
    path = tmp_path / "synthetic.docx"
    document.save(str(path))
    return path


@pytest.fixture()
def synthetic_xlsx(tmp_path: Path) -> Path:
    import openpyxl

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Sheet1"
    sheet["A1"] = "header"
    workbook.properties.creator = "Synthetic Creator"
    workbook.properties.language = "en-GB"
    workbook.properties.lastModifiedBy = "Synthetic XLSX Modifier"
    workbook.properties.modified = datetime(2026, 9, 1, 12, 0, 0)
    path = tmp_path / "synthetic.xlsx"
    workbook.save(str(path))
    return path


def test_docx_extractor_reads_author_as_prepared_by(synthetic_docx: Path) -> None:
    metadata = extract_asset_metadata_from_docx_properties(synthetic_docx)
    assert metadata.roles.prepared_by.raw_value == "Synthetic Author"
    assert metadata.roles.prepared_by.source == MetadataSource.FILE_PROPERTIES


def test_docx_extractor_reads_last_modified_by_and_date(synthetic_docx: Path) -> None:
    metadata = extract_asset_metadata_from_docx_properties(synthetic_docx)
    assert metadata.audit.last_modified_by.raw_value == "Synthetic Modifier"
    assert metadata.audit.last_modified_date.normalized_value is not None
    assert metadata.audit.last_modified_date.normalized_value.isoformat() == "2026-08-14"


def test_docx_extractor_reads_language(synthetic_docx: Path) -> None:
    metadata = extract_asset_metadata_from_docx_properties(synthetic_docx)
    assert metadata.identity.language_code.raw_value == "en-US"


def test_docx_extractor_file_format_from_filename_never_content(synthetic_docx: Path) -> None:
    metadata = extract_asset_metadata_from_docx_properties(synthetic_docx)
    assert metadata.identity.file_format.raw_value == "docx"
    assert metadata.identity.file_format.source == MetadataSource.FILENAME


def test_docx_extractor_never_modifies_the_source_file(synthetic_docx: Path) -> None:
    before = _sha256(synthetic_docx)
    extract_asset_metadata_from_docx_properties(synthetic_docx)
    after = _sha256(synthetic_docx)
    assert before == after


def test_docx_extractor_never_maps_numeric_revision_save_count() -> None:
    """OOXML core_properties.revision is a bare save-count integer, not a
    document revision label -- this extractor must never read or expose
    it anywhere in its own output. Verified structurally: no field in
    KnowledgeAssetMetadata is populated from it (the extractor source
    itself never references `.revision` at all -- see the module's own
    docstring for the rationale).
    """
    assert _revision_attribute_accesses_in_module() == []


def test_docx_extractor_missing_property_yields_none_not_a_guess(tmp_path: Path) -> None:
    from docx import Document

    document = Document()
    document.add_paragraph("No properties explicitly set.")
    path = tmp_path / "bare.docx"
    document.save(str(path))

    metadata = extract_asset_metadata_from_docx_properties(path)
    assert metadata.identity.language_code.raw_value is None


def test_xlsx_extractor_reads_creator_as_prepared_by(synthetic_xlsx: Path) -> None:
    metadata = extract_asset_metadata_from_xlsx_properties(synthetic_xlsx)
    assert metadata.roles.prepared_by.raw_value == "Synthetic Creator"
    assert metadata.roles.prepared_by.source == MetadataSource.FILE_PROPERTIES


def test_xlsx_extractor_reads_last_modified_by_and_date(synthetic_xlsx: Path) -> None:
    """`openpyxl.Workbook.save()` itself re-stamps `properties.modified`
    to the real current save time regardless of what a caller assigned
    beforehand (confirmed empirically while writing this test -- an
    openpyxl behavior, not an extractor defect) -- so this test proves
    self-consistency (whatever real value ends up in the saved file's own
    `docProps/core.xml` is what the extractor reads and correctly
    normalizes) rather than asserting a specific fixed date. Excel-serial/
    ISO-8601 normalization correctness itself is already exhaustively
    covered by test_asset_metadata.py's own dedicated date-normalization
    tests.
    """
    import openpyxl

    reloaded = openpyxl.load_workbook(str(synthetic_xlsx), read_only=True)
    try:
        actual_modified = reloaded.properties.modified
    finally:
        reloaded.close()

    metadata = extract_asset_metadata_from_xlsx_properties(synthetic_xlsx)
    assert metadata.audit.last_modified_by.raw_value == "Synthetic XLSX Modifier"
    assert metadata.audit.last_modified_date.normalized_value is not None
    assert metadata.audit.last_modified_date.normalized_value == actual_modified.date()


def test_xlsx_extractor_reads_language(synthetic_xlsx: Path) -> None:
    metadata = extract_asset_metadata_from_xlsx_properties(synthetic_xlsx)
    assert metadata.identity.language_code.raw_value == "en-GB"


def test_xlsx_extractor_file_format_from_filename(synthetic_xlsx: Path) -> None:
    metadata = extract_asset_metadata_from_xlsx_properties(synthetic_xlsx)
    assert metadata.identity.file_format.raw_value == "xlsx"
    assert metadata.identity.file_format.source == MetadataSource.FILENAME


def test_xlsx_extractor_never_modifies_the_source_file(synthetic_xlsx: Path) -> None:
    before = _sha256(synthetic_xlsx)
    extract_asset_metadata_from_xlsx_properties(synthetic_xlsx)
    after = _sha256(synthetic_xlsx)
    assert before == after


def test_xlsx_extractor_never_maps_numeric_revision_save_count() -> None:
    assert _revision_attribute_accesses_in_module() == []
