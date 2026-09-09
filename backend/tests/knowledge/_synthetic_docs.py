"""Minimal SYNTHETIC document builders for A5 ingestion tests -- every
value here is invented and non-sensitive (A5 instruction section 48).
Never imports or embeds any real corpus content. Not a test module
itself (no `test_*` functions) -- imported by the real test files.
"""
from __future__ import annotations

import io
import struct
import zipfile

_CFB_SECTOR_SIZE = 512
_CFB_FREESECT = 0xFFFFFFFF
_CFB_ENDOFCHAIN = 0xFFFFFFFE
_CFB_FATSECT = 0xFFFFFFFD


def make_minimal_png(color: tuple[int, int, int] = (200, 30, 30)) -> bytes:
    from PIL import Image

    buffer = io.BytesIO()
    Image.new("RGB", (4, 4), color=color).save(buffer, format="PNG")
    return buffer.getvalue()


def make_minimal_pdf(text: str = "Synthetic PDF page") -> bytes:
    """A hand-assembled, minimally-valid single-page PDF with a real
    extractable text content stream and a correct xref table -- built
    without any new dependency (no reportlab).
    """
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /Resources << /Font << /F1 4 0 R >> >> /MediaBox [0 0 200 200] /Contents 5 0 R >>",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    stream_body = f"BT /F1 24 Tf 10 100 Td ({text}) Tj ET".encode("latin-1")
    objects.append(b"<< /Length %d >>\nstream\n" % len(stream_body) + stream_body + b"\nendstream")

    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for index, body in enumerate(objects, start=1):
        offsets.append(out.tell())
        out.write(f"{index} 0 obj\n".encode())
        out.write(body)
        out.write(b"\nendobj\n")
    xref_offset = out.tell()
    count = len(objects) + 1
    out.write(f"xref\n0 {count}\n".encode())
    out.write(b"0000000000 65535 f \n")
    for offset in offsets:
        out.write(f"{offset:010d} 00000 n \n".encode())
    out.write(f"trailer\n<< /Size {count} /Root 1 0 R >>\nstartxref\n{xref_offset}\n%%EOF".encode())
    return out.getvalue()


def make_minimal_xlsx(sheets: dict[str, list[list[object]]]) -> bytes:
    """`sheets`: sheet name -> list of rows (each row a list of cell
    values; the first row is treated as headers by the extractor).
    A row value that is a string starting with "=" is written as a real
    formula cell (never evaluated by this codebase).
    """
    import openpyxl

    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)
    for sheet_name, rows in sheets.items():
        worksheet = workbook.create_sheet(sheet_name)
        for row in rows:
            worksheet.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def make_minimal_docx(
    *,
    heading: str = "Synthetic Procedure",
    paragraphs: list[str] | None = None,
    table_rows: list[list[str]] | None = None,
    image_png: bytes | None = None,
    hyperlink_text: str | None = None,
    hyperlink_url: str | None = None,
) -> bytes:
    from docx import Document

    document = Document()
    document.add_heading(heading, level=1)
    for paragraph_text in paragraphs or ["Invented, non-sensitive body text."]:
        document.add_paragraph(paragraph_text)
    if table_rows:
        table = document.add_table(rows=len(table_rows), cols=len(table_rows[0]))
        for row_index, row_values in enumerate(table_rows):
            for col_index, value in enumerate(row_values):
                table.cell(row_index, col_index).text = value
    if image_png:
        document.add_picture(io.BytesIO(image_png))
    if hyperlink_text and hyperlink_url:
        # python-docx has no simple public "add hyperlink" API pre-1.2's
        # own limited support; append via the low-level OOXML relationship
        # mechanism the library itself provides for this exact purpose.
        part = document.part
        r_id = part.relate_to(hyperlink_url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
        paragraph = document.add_paragraph()
        hyperlink = paragraph._p.makeelement(
            "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}hyperlink",
            {"{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id": r_id},
        )
        run = paragraph.add_run(hyperlink_text)
        run_element = run._r
        paragraph._p.remove(run_element)
        hyperlink.append(run_element)
        paragraph._p.append(hyperlink)

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def inject_embedded_member(docx_bytes: bytes, member_path: str, member_bytes: bytes) -> bytes:
    """Post-processes a DOCX zip package to add one raw member (e.g.
    "word/embeddings/nested.xlsx") -- exercises the recursive-extraction
    fallback discovery path (raw `word/embeddings/`/`word/media/` scan)
    without needing to hand-construct real OOXML relationship XML.
    """
    input_buffer = io.BytesIO(docx_bytes)
    output_buffer = io.BytesIO()
    with zipfile.ZipFile(input_buffer) as source, zipfile.ZipFile(output_buffer, "w") as destination:
        for item in source.infolist():
            destination.writestr(item, source.read(item.filename))
        destination.writestr(member_path, member_bytes)
    return output_buffer.getvalue()


def _cfb_directory_entry(name: str, obj_type: int, *, left: int, right: int, child: int, start_sector: int, size: int) -> bytes:
    name_utf16 = name.encode("utf-16-le") + b"\x00\x00"
    name_length = len(name_utf16)
    name_field = name_utf16.ljust(64, b"\x00")[:64]
    entry = name_field
    entry += struct.pack("<H", name_length)
    entry += struct.pack("<B", obj_type)
    entry += struct.pack("<B", 1)  # color flag: black
    entry += struct.pack("<I", left)
    entry += struct.pack("<I", right)
    entry += struct.pack("<I", child)
    entry += b"\x00" * 16  # CLSID
    entry += b"\x00" * 4  # state bits
    entry += b"\x00" * 8  # creation time
    entry += b"\x00" * 8  # modified time
    entry += struct.pack("<I", start_sector)
    entry += struct.pack("<Q", size)
    assert len(entry) == 128
    return entry


def make_minimal_cfb_with_stream(stream_name: str, stream_data: bytes) -> bytes:
    """Hand-built, MINIMAL, structurally-valid OLE2 Compound-File-Binary
    container holding exactly one named stream -- built without any
    write-capable OLE library (`olefile` is READ-ONLY by design), so
    this is the only way to produce a real, `olefile`-openable
    synthetic CFB fixture for tests. `stream_data` is kept >= 4096
    bytes by convention at call sites so it always uses regular (never
    mini-stream) sectors, avoiding the CFB mini-FAT path entirely --
    the smallest correct layout: one FAT sector, one directory sector
    (Root Entry + the one named stream), then the stream's own data
    sectors. Validated by direct `olefile.OleFileIO` round-trip in
    `test_ingestion_extractors_ole.py`.
    """
    num_data_sectors = max(1, (len(stream_data) + _CFB_SECTOR_SIZE - 1) // _CFB_SECTOR_SIZE)
    padded_data = stream_data + b"\x00" * (num_data_sectors * _CFB_SECTOR_SIZE - len(stream_data))

    fat_entries = [_CFB_FREESECT] * (_CFB_SECTOR_SIZE // 4)
    fat_entries[0] = _CFB_FATSECT
    fat_entries[1] = _CFB_ENDOFCHAIN
    for i in range(num_data_sectors):
        fat_entries[2 + i] = (2 + i + 1) if i < num_data_sectors - 1 else _CFB_ENDOFCHAIN
    fat_sector = b"".join(struct.pack("<I", v) for v in fat_entries)

    root_entry = _cfb_directory_entry("Root Entry", 5, left=_CFB_FREESECT, right=_CFB_FREESECT, child=1, start_sector=_CFB_ENDOFCHAIN, size=0)
    stream_entry = _cfb_directory_entry(
        stream_name, 2, left=_CFB_FREESECT, right=_CFB_FREESECT, child=_CFB_FREESECT, start_sector=2, size=len(stream_data)
    )
    directory_sector = root_entry + stream_entry + b"\x00" * 128 + b"\x00" * 128

    header = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"
    header += b"\x00" * 16
    header += struct.pack("<H", 0x003E)
    header += struct.pack("<H", 0x0003)
    header += struct.pack("<H", 0xFFFE)
    header += struct.pack("<H", 9)
    header += struct.pack("<H", 6)
    header += b"\x00" * 6
    header += struct.pack("<I", 0)
    header += struct.pack("<I", 1)
    header += struct.pack("<I", 1)
    header += struct.pack("<I", 0)
    header += struct.pack("<I", 4096)
    header += struct.pack("<I", _CFB_ENDOFCHAIN)
    header += struct.pack("<I", 0)
    header += struct.pack("<I", _CFB_ENDOFCHAIN)
    header += struct.pack("<I", 0)
    difat = [0] + [_CFB_FREESECT] * 108
    header += b"".join(struct.pack("<I", v) for v in difat)
    assert len(header) == _CFB_SECTOR_SIZE

    return header + fat_sector + directory_sector + padded_data


def make_ole10_native_stream(filename: str, payload: bytes) -> bytes:
    """Builds the raw bytes of a classic OLE 1.0 "Package" `"\\x01Ole10Native"`
    stream (4-byte total size, 2-byte version, NUL-terminated display
    filename, NUL-terminated source path, an 8-byte reserved block,
    NUL-terminated temp path, 4-byte payload length, then the raw
    payload bytes) -- the exact layout
    `backend/knowledge/ingestion/extractors/ole.py`'s
    `_parse_ole10_native_package` expects, empirically validated against
    the real corpus.
    """
    version = struct.pack("<H", 2)
    filename_field = filename.encode("latin-1") + b"\x00"
    source_path_field = f"C:\\synthetic\\{filename}".encode("latin-1") + b"\x00"
    reserved = b"\x00" * 8
    temp_path_field = f"C:\\synthetic\\temp\\{filename}".encode("latin-1") + b"\x00"
    data_length_field = struct.pack("<I", len(payload))

    body = version + filename_field + source_path_field + reserved + temp_path_field + data_length_field + payload
    total_size = struct.pack("<I", len(body))
    return total_size + body


def make_ole_package_object(filename: str, payload: bytes) -> bytes:
    """A complete, `olefile`-openable OLE2 CFB container wrapping one
    `"\\x01Ole10Native"` Package stream for `filename`/`payload` -- the
    synthetic equivalent of a real Word "Insert Object > Create from
    File" embed. `payload` is padded internally (via the stream's own
    size accounting) if needed to keep the overall stream >= 4096 bytes,
    so tests never need to supply an artificially large fixture just to
    stay off the CFB mini-stream path.
    """
    stream = make_ole10_native_stream(filename, payload)
    if len(stream) < 4096:
        stream = stream + b"\x00" * (4096 - len(stream))
    return make_minimal_cfb_with_stream("\x01Ole10Native", stream)
